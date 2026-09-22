from pathlib import Path
import re, math
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
OUT = HERE / 'One DHS Program Automations and Workflows.docx'
NAVY = '203D53'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(.66), Inches(.65)
sec.left_margin, sec.right_margin = Inches(.72), Inches(.72)
sec.header_distance, sec.footer_distance = Inches(.25), Inches(.28)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'Caption']:
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st.font.color.rgb = RGBColor(0,0,0)
    st.paragraph_format.space_after = Pt(6)
for st in doc.styles:
    for border in st.element.findall('.//' + qn('w:pBdr')):
        border.getparent().remove(border)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing = 1.09
doc.styles['Title'].font.size = Pt(27)
doc.styles['Title'].font.bold = True
doc.styles['Subtitle'].font.size = Pt(16)
doc.styles['Heading 1'].font.size = Pt(19)
doc.styles['Heading 2'].font.size = Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(9)
doc.styles['Caption'].font.size = Pt(10)
header = sec.header.paragraphs[0]
header.text = 'ONE DHS PEOPLE ACCESS AND CULTURE PROGRAM'
header.runs[0].font.size = Pt(8)
foot = sec.footer.paragraphs[0]
foot.text = 'Consultant reference  |  September 6 2026'
foot.runs[0].font.size = Pt(8)
foot.add_run(' '*8 + 'Page ').font.size = Pt(8)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); foot._p.append(fld)
doc.core_properties.title = 'One DHS Program Automations and Workflows'
doc.core_properties.subject = 'Implementation review and workflow reference for Gary Banks'
doc.core_properties.author = 'Prepared for Gary Banks'

def p(text, style=None):
    return doc.add_paragraph(text, style)
def h(text): doc.add_heading(text,2)
def page(title):
    doc.add_page_break(); doc.add_heading(title,1)
def ref(text):
    r=p('Implementation references  '+text)
    for run in r.runs: run.font.size=Pt(9); run.font.color.rgb=RGBColor.from_string('555555')
def table(headers, rows, widths=None, size=10):
    t=doc.add_table(rows=1, cols=len(headers)); t.autofit=False
    if widths:
        for c,w in zip(t.columns,widths): c.width=Inches(w)
    for c,text in zip(t.rows[0].cells,headers): c.text=text
    repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for c,text in zip(t.add_row().cells,row): c.text=str(text)
    for i,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr(); cant=OxmlElement('w:cantSplit'); trpr.append(cant)
        for j,c in enumerate(row.cells):
            if widths:c.width=Inches(widths[j])
            c.vertical_alignment=1
            pr=c._tc.get_or_add_tcPr()
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),NAVY if i==0 else ('F1F4F6' if i%2 else 'FFFFFF'));pr.append(sh)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'CED5DB');borders.append(e)
            pr.append(borders)
            mar=OxmlElement('w:tcMar')
            for edge,val in [('top','90'),('bottom','90'),('left','100'),('right','100')]:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),val);e.set(qn('w:type'),'dxa');mar.append(e)
            pr.append(mar)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(2);para.paragraph_format.line_spacing=1.04
                for r in para.runs:
                    r.font.size=Pt(size)
                    if i==0:r.bold=True;r.font.color.rgb=RGBColor(255,255,255)
    p('') .paragraph_format.space_after=Pt(0)
    return t

FONT=ImageFont.truetype('C:/Windows/Fonts/calibri.ttf',38)
BOLD=ImageFont.truetype('C:/Windows/Fonts/calibrib.ttf',38)
def diagram(name,nodes,edges,caption):
    im=Image.new('RGB',(1680,650),'white');d=ImageDraw.Draw(im)
    def arrow(a,b):
        x1,y1,x2,y2=a+b
        d.line((x1,y1,x2,y2),fill='#647786',width=5)
        ang=math.atan2(y2-y1,x2-x1);L=17
        d.polygon([(x2,y2),(x2-L*math.cos(ang-.5),y2-L*math.sin(ang-.5)),(x2-L*math.cos(ang+.5),y2-L*math.sin(ang+.5))],fill='#647786')
    for a,b in edges:arrow(a,b)
    for x,y,w,hh,txt in nodes:
        d.rounded_rectangle((x,y,x+w,y+hh),radius=12,fill='#F0F4F6',outline='#73899A',width=3)
        words=txt.split();lines=[];line=''
        for word in words:
            test=(line+' '+word).strip()
            if d.textlength(test,font=FONT)>w-38 and line:lines.append(line);line=word
            else:line=test
        if line:lines.append(line)
        yy=y+(hh-len(lines)*45)/2
        for line in lines:
            d.text((x+w/2,yy),line,fill='#172E40',font=FONT,anchor='mt');yy+=45
    path=HERE/(name+'.png');im.save(path)
    shape=doc.add_picture(str(path),width=Inches(7.02))
    shape._inline.docPr.set('descr',caption)
    p(caption,'Caption')

p('One DHS People Access and Culture Program','Title')
p('Automations and workflows','Subtitle')
p('Prepared for Gary Banks\nConsultant reference document\nImplementation snapshot September 6 2026')
h('The main finding')
p('The program has a substantial workflow foundation, but it is not yet a fully connected, autonomous operating system. It combines working page editing, scoped learning and resource discovery, deterministic assistance, owner controls, and a developing collaboration workspace. Several important capabilities exist in code but are not enabled in the inspected local configuration. Others remain design commitments.')
p('This document explains the program as one connected ecosystem. It covers staff journeys, resource handling, consultation, agent activity, publishing, collaboration, safety controls, and maintenance. Diagrams show how the parts relate; the register at the end accounts for every named tool in the current agent catalog.')
h('How to read the findings')
table(['Status','Meaning'],[
('Observed','A recorded local end to end check supports the behavior. This does not certify the live deployment.'),
('Implemented','A concrete code path exists. It may still need configuration, data, or end to end testing.'),
('Conditional','The code requires an explicit switch, credential, approved setup, or other prerequisite.'),
('Declared or planned','A description, schema, catalog entry, or program intention exists, but that is not evidence of an operating automation.'),
('Human led','People initiate or carry out the work. The application may supply materials or record decisions.')],[1.25,5.75])
p('Scope of evidence: the current local repository, configuration indicators, routes, workflow code, and saved verification notes. No paid research was run, no agency system was accessed, and no live deployment or Teams synchronization was certified for this report. No application settings were changed.')

page('1 How the program fits together')
diagram('ecosystem',[
(20,30,470,155,'Staff choose their work context'),(610,30,470,155,'Learning resources and Ask'),(1190,30,470,155,'Use guidance in everyday work'),
(20,390,470,170,'Amplify Equity offers voluntary engagement'),(610,390,470,170,'One DSD Team examines shared questions'),(1190,390,470,170,'Consultant and responsible leaders decide')],
[((490,108),(610,108)),((1080,108),(1190,108)),((255,185),(255,390)),((490,475),(610,475)),((1080,475),(1190,475)),((1425,390),(1425,185))],
'Figure 1  Learning and action connect with staff engagement and structural review. The lower pathway is human led; it is not an automatic transfer of staff conversations.')
p('The shared foundation is One DHS. One DSD supplies division specific material and practice without requiring a separate technical installation. A staff member can explore learning and resources, ask a question, and follow an appropriate next step. The One DSD Team and Amplify Equity connect to this foundation in different ways.')
p('A work context is not a security permission. The current One DHS and One DSD selection changes relevance and scope. It does not establish that a visitor holds a particular job, belongs to a division, or has authority to edit its content. The broader administration by administration permission model still needs implementation and verification.')
h('Connections that must remain explicit')
p('Learning can support a workplace task without becoming mandatory training. Engagement can surface a structural question without turning participants into investigators. A consultation can support a decision without replacing the person authorized to make it. External research can inform the work without acquiring the status of official DHS guidance.')
ref('S01 S02 S09 S10')

page('2 Learning and resources')
p('Trigger: a person opens Learning, chooses a theme, searches, or selects a work role and task. Output: relevant published resources and optional next steps. This is implemented discovery and recommendation logic, not a fully individualized course generation service.')
table(['Workflow','Current behavior and limit'],[
('Thematic discovery','One resource can appear under several themes. References resolve against the scoped published collection, preventing a recovered candidate from appearing merely because a theme names it.'),
('Course entry','Course tiles provide a title, description, cover, and route to the learning resource. Existing Library and Resources routes remain available.'),
('Learning for your work','Ten work profiles and four task choices supply curated resource identifiers. Task matches are prioritized, duplicates removed, and missing or unpublished items excluded.'),
('Learning sequence','Six conceptual stages support orientation through leadership and continuity. They guide learning rather than require staff to unlock resources in order.'),
('Next practice','The graduation helper recommends a next independent action. Its presence does not establish course completion, certification, or an official learning record.')],[1.48,5.52])
h('What is not yet demonstrated')
p('The job map does not yet cause Ask to construct a position specific curriculum from a complete organizational graph. Local progress tools appear in the catalog, but this review did not establish a complete persistent progress, bookmark, or journal workflow in the current app. Those features must not be described as finished solely because they exist in another repository or in a tool description.')
p('Each lesson should have three to five observable objectives. That is the program authoring standard, not evidence that every recovered lesson currently complies. A future review should check objectives, examples, accessible media, and a useful application task while preserving immediate access to the resource.')
h('Workforce relevance')
p('The workforce map provides an editable organizational baseline and role descriptions. It is not a live HR directory, a funded position inventory, or a verified record of every current job. Responsible leaders need a way to validate changes in their own area. No automatic import from DHS personnel systems is present.')
ref('S01 S02 S03')

page('3 Ask and research')
diagram('ask',[
(20,30,460,160,'Question and optional context'),(610,30,460,160,'Safety checks and scoped retrieval'),(1200,30,460,160,'Answer with sources and limits'),
(20,395,460,160,'Public research selected or needed'),(610,395,460,160,'Owner switch and provider readiness'),(1200,395,460,160,'External evidence kept distinct')],
[((480,110),(610,110)),((1070,110),(1200,110)),((840,190),(250,395)),((480,475),(610,475)),((1070,475),(1200,475)),((1430,395),(1430,190))],
'Figure 2  Ask uses the published program collection first. Public research is conditional. If a provider is unavailable, the program must not imply that fresh research occurred.')
p('Ask classifies the question, applies privacy and misuse checks, searches the selected published collection, and composes a response. It identifies gaps and attaches available source references. It can suggest independent practice or an appropriate contact; consultation is not the required destination for every question.')
p('The function named semantic retrieval currently uses text scoring and an intent boost. It is not evidence of embedding based retrieval or knowledge graph traversal. The present workflow is therefore more limited than the intended adaptive research and knowledge system.')
h('Three research experiences')
p('Ranked web results, a current answer, and deeper research are three experiences within the research layer. The current provider code selects the Perplexity Agent route or a test fixture; these are not three independently connected research vendors. Legacy provider selections are normalized to the newer route.')
p('The inspected local configuration has no Perplexity credential and no saved research policy enabling the service. External research is therefore not demonstrated as available here. The code screens the outgoing query and records usage metadata rather than treating an entire consultation packet as research input. Staff access is not divided into research tiers; the consultant controls enablement.')
ref('S04 S05')

page('4 Agents and actual model behavior')
p('An agent entry states a role and permitted tools. It is not proof of an independent worker running in the background. The current registry has ten agents. Most behavior is directed by explicit application code; it is not a general network of agents freely discovering and executing every program task.')
table(['Agent','What exists now'],[
('Program orchestrator','Runs a bounded cycle and records findings, triage changes, flags, and proposals when conditions allow.'),
('Ask concierge','Handles scoped questions and optional research through the Ask path.'),
('Cultural intelligence guide','Supports brief retrieval, progressive detail, limits, and relevant next steps.'),
('Librarian','Provides classification and stale content checks. A recommendation is not an automatic publication decision.'),
('Accessibility reviewer','Produces review hints and checklists. It does not certify WCAG compliance.'),
('Consultation intake','Validates and prepares requests. Submission is conditional on operational readiness.'),
('Graduation coach','Suggests independent next practice. It does not award official completion.'),
('Embed advisor','Provides question sets and practice prompts; no separate autonomous entry point was established.'),
('Content sentinel','Registered, with its feature flag now enforced. The separate cycle librarian review does not imply this agent is enabled.'),
('Evaluation steward','Supports evaluation and proposals. Automatic model comparison remains disabled.')],[1.55,5.45],10)
h('Model readiness')
p('The code includes a working Anthropic provider adapter, but the inspected local configuration lacks its credential. Deterministic composition remains a real fallback; it must not be presented as live model reasoning. The OpenAI provider is explicitly a stub that cannot generate responses. Switching to a future model is not automatic.')
p('A model upgrade needs an implemented adapter, credentials, supported input and output contracts, evaluations, a release decision, and a recovery path. The intended ability to evolve is sensible; current registry flexibility alone does not deliver it.')
ref('S04 S06 S07')

page('5 The recurring review cycle')
diagram('cycle',[
(20,25,470,170,'Owner runs a cycle or scheduler calls it'),(605,25,470,170,'Read policy and apply retention duties'),(1190,25,470,170,'Inspect eligible work and content'),
(20,395,470,170,'Owner reviews results and proposals'),(605,395,470,170,'Save report and decision records'),(1190,395,470,170,'Apply permitted bounded changes')],
[((490,110),(605,110)),((1075,110),(1190,110)),((1425,195),(1425,395)),((1190,480),(1075,480)),((605,480),(490,480))],
'Figure 3  The review cycle has defined steps. A saved report is evidence of a run; a schedule in a configuration file is only evidence that scheduling was declared.')
p('The declared Vercel schedule calls the orchestration route daily at 12 UTC. The route requires its scheduler secret. The local configuration contains that secret, but a development server does not independently run the Vercel schedule. No successful live scheduled run was verified in this review.')
p('The cycle checks retention first, then examines eligible consultation work, stale resources, brief quality, and accessibility concerns. At sufficient authority it can refresh preparation packets, move received requests to under review, and add a priority pin or reversible review flag. It excludes ineligible or expired requests and uses version checks for consultation changes.')
p('At the highest configured level it can create proposals for the owner, such as reviewing content or changing capacity or policy settings. A proposal is not a self approved program change. A generative summary is conditional on provider readiness.')
h('Important limitations')
p('Resource currency and accessibility checks now read the selected published collection, including shared and DSD resources. The brief quality check still reads the separate brief definitions. Cycle undo clears eligible automation owned pins and flags; it does not reverse consultation status history or restore deleted information.')
ref('S07 S08')

page('6 Authority safety and recovery')
table(['Control','Behavior and qualification'],[
('Autonomy ceiling','The runtime combines the program ceiling, agent limits, and a hard maximum. The inspected local environment sets A5; that alone does not enable missing providers or integrations.'),
('Tool admission','Registered and enabled tool, agent allowlist, policy allowance, required authority, owner restriction, and dry run behavior are checked on the shared tool path.'),
('Owner stop','Stops agent work through policy. Retention and requester privacy rights have deliberate exceptions. A short policy cache means this is not a guarantee that all activity halts at the same instant.'),
('Human decisions','Publishing and owner decisions use specific protected routes. Catalog phrases such as human approval or preview are not, by themselves, an implemented approval screen for every tool.'),
('Audit and memory','Records decisions and workflow outcomes with redaction. This is not a license to retain private staff conversations or personal case information.'),
('Recovery','Content versions support restoration. Consultation version checks prevent stale updates. Cycle undo has narrower limits. Data deletion is not generally reversible.')],[1.4,5.6])
h('Corrections made during this review')
p('The cycle now respects orchestrator disablement at lower authority levels. Brief and accessibility checks pass through their assigned agent controls. Agent feature flags, roles, environments, and the stale flag writing switch are enforced. Generated proposals now pass through the same A5 check as other proposals, including a policy check after the provider returns.')
p('Model readiness now checks whether the registered model is callable, not just whether a flag and credential exist. Further verification is still needed across the whole application. A named tool is not necessarily backed by a separate service, and a timeout in catalog metadata is not proof of enforced cancellation.')
h('Recommended proof before wider operation')
p('Run the same scenario with authority lowered, the relevant agent disabled, the owner stop active, a duplicate request, a stale version, and a provider failure. Confirm both the response and the absence of unintended writes. Keep retention exceptions explicit and test them separately. Do not describe these recommendations as checks already passed.')
ref('S06 S07 S08 S13')

page('7 Consultation from preparation to closure')
diagram('consultation',[
(20,25,470,175,'Staff prepare a general work request'),(605,25,470,175,'Readiness checks allow submission'),(1190,25,470,175,'Save request and private access key'),
(20,390,470,175,'Close withdraw or apply retention'),(605,390,470,175,'Consultant reviews and follows up'),(1190,390,470,175,'Confirm eligibility and prepare queue')],
[((490,112),(605,112)),((1075,112),(1190,112)),((1425,200),(1425,390)),((1190,477),(1075,477)),((605,477),(490,477))],
'Figure 4  Preparation and submission are different steps. The current local intake switch is off. No calendar invitation or Teams message is sent by this flow.')
p('Preparation can organize a question and draft a useful agenda without accepting a real service request. Submission has stronger prerequisites: the enable switch, durable operational storage, tracking and rate limit secrets, approved activation evidence, retention arrangements, and readiness for requester rights and service delivery.')
p('When enabled, a request includes confirmation of what may be shared. Repeated submissions are handled through an idempotency key to reduce duplicates. DSD eligibility is confirmed rather than inferred. The requester receives limited access to their own record, while owner notes and internal preparation remain separate.')
p('The requester can view status, correct permitted fields, rotate the access key, and withdraw before scheduling under the defined rules. Version checks prevent a correction from silently overwriting newer work. These rights remain available under the owner stop when their required tracking configuration is ready.')
p('The consultant controls the service outcome and status changes. Calendar handoff text is a draft to copy, not a scheduling integration. Expired records are redacted or removed according to the configured retention lifecycle. A deployed scheduler and operational evidence are needed before claiming reliable unattended retention.')
ref('S08 S11')

page('8 Editing review and publication')
diagram('editing',[
(20,30,470,160,'Owner edits a registered content field'),(605,30,470,160,'Save draft without changing staff copy'),(1190,30,470,160,'Record required reviews'),
(20,395,470,160,'Restore an earlier version if needed'),(605,395,470,160,'Read the released version on the page'),(1190,395,470,160,'Owner publishes within allowed scope')],
[((490,110),(605,110)),((1075,110),(1190,110)),((1425,190),(1425,395)),((1190,475),(1075,475)),((605,475),(490,475))],
'Figure 5  Drafts stay separate from displayed content. The saved local verification demonstrates draft persistence, publication after five reviews, and restoration on a private workforce page.')
p('The current editing registry contains 76 areas. Local selection enables 62; fourteen community brief areas remain on their previous presentation pending separate reviews. The saved verification note records a successful draft, reload, review, publication, reload, and restoration test. This is evidence for that tested local path, not a claim that every page and field was exercised.')
p('Editing uses owner authentication and protected mutations. Existing scope and release rules remain in place. Page content can be held as a draft and restored through recorded versions. The DHS logo remains outside editable fields.')
h('The boundary of edit anywhere')
p('The current implementation does not establish unlimited inline editing of every layout, resource file, media object, link block, or community brief. Registered fields and enabled areas define what is editable. Resource draft and release routes are separate from the generalized page content editor; database setup for one does not automatically activate the other.')
p('This report does not recommend bypassing review to deliver broader editability. The next proof should cover representative public pages, resource records, complex fields, failed reviews, conflicting edits, and recovery. The live Vercel app was not deployed as part of the recorded local editing work.')
ref('S12 S13')

page('9 Resource ingestion and the knowledge graph')
diagram('corpus',[
(20,25,470,175,'Collect source files links and records'),(605,25,470,175,'Identify duplicates and retain receipts'),(1190,25,470,175,'Stage assets families and relationships'),
(20,395,470,175,'Discover through published scope'),(605,395,470,175,'Release the selected material'),(1190,395,470,175,'Review content access and language')],
[((490,112),(605,112)),((1075,112),(1190,112)),((1425,200),(1425,395)),((1190,482),(1075,482)),((605,482),(490,482))],
'Figure 6  A retained source record is not the same as a usable published resource. Staging, review, release, and staff discovery are distinct steps.')
p('The repository includes operator run scripts to capture sources, verify the ledger, build and load staged imports, form canonical projections, store source objects, release seed content, and benchmark batches. These are implemented processing tools, not evidence of an unattended watcher that continually ingests every desktop folder.')
p('The historical inventory of 524 raw records and 557 disposition receipts describes different units. Families, child resources, merged records, internal material, audio, and course assets cannot be added together as if each were one distinct staff resource. Preserve their relationships and the seed items rather than discard them to force one count.')
p('The local content source remains static. Consequently, database assets or recovered candidates do not automatically become visible in Learning or available to Ask. Publication selection and the runtime content source must both be correct. This is a material explanation for a much smaller visible resource count.')
h('Graph engineering and scale')
p('Relationship data and graph preparation work are part of the foundation. The inspected Ask retrieval path does not traverse that graph. Likewise, scripts and a benchmark command do not prove successful operation at 30,000 to 50,000 resources. That requires measured import, retrieval, release, duplicate handling, recovery, and accessibility tests with a representative collection.')
ref('S14 S04 S12')

page('10 The One DSD Team and Amplify Equity')
diagram('teams',[
(20,25,470,175,'Amplify participants choose topics'),(605,25,470,175,'Co leads carry agreed shared questions'),(1190,25,470,175,'One DSD Team examines structural needs'),
(20,395,470,175,'Staff receive a clear response'),(605,395,470,175,'Review results and decide next steps'),(1190,395,470,175,'Responsible owner takes up the work')],
[((490,112),(605,112)),((1075,112),(1190,112)),((1425,200),(1425,395)),((1190,482),(1075,482)),((605,482),(490,482))],
'Figure 7  This is a proposed continuing human workflow, supported by program materials. Questions move with agreement, not through automatic monitoring of staff conversations.')
p('The One DSD Team is the established equity committee described by the owner. It supports the consultant with varied work expertise, review, and follow through. Amplify Equity is a prepared, voluntary engagement space with one or two staff co leads. It should not be represented as already relaunched or assigned a confirmed meeting schedule.')
p('Amplify can hold conversation, cultural learning, wellness and workplace topics, and staff perspectives on fairness and access. Co leads provide the connection to the One DSD Team. The consultant is available by invitation, not the routine facilitator or manager of the group. Resources should support staff facilitation without making every gathering dependent on the consultant.')
h('What the application currently supports')
p('The private One DSD workspace implements changes to summaries, channel purposes, threads, replies, posts, meeting notes, action items, and feedback. Mutations currently run through owner access. The data model includes broader membership, attachment, notification, and automation concepts, but their presence does not establish a complete multi member collaboration service.')
p('There is no demonstrated automation that converts Amplify discussion into structural assignments. That boundary should remain intentional: agree on the question, omit personal details, identify a responsible person, and return a response. This supports continuous improvement without turning the engagement space into surveillance.')
ref('S09 S10')

page('11 Microsoft Teams and other connections')
table(['Connection','What it does and does not do'],[
('Personal Microsoft Teams','Program links can open the prepared communities. Opening a destination is navigation, not synchronized membership, files, conversations, or workflow state.'),
('Teams bridge','The current bridge module explicitly reports off, disconnected, no external reads, and no external writes. It cannot authenticate or send a message.'),
('Official agency Teams','No official tenant integration is part of this inspected implementation. Future transfer of a template is a separate, authorized activity, not an automatic migration.'),
('Native collaboration','The local program workspace is its own record. A Teams post does not automatically update it, and a local action item does not become a Teams task.'),
('Calendar and messages','Calendar text can be drafted. Scheduling, email delivery, Teams posting, and reminders are not established integrations.'),
('Database','Saved page editing uses its selected database areas. Operational storage and resource publication have separate configuration; one working database path does not activate them all.')],[1.53,5.47])
h('A practical future connection')
p('The code records an attended personal Teams pilot as the selected future bridge approach. Before such a pilot, verify the account and destination, use synthetic material, prevent duplicate sends, and review each outbound action. A durable independent Microsoft integration would require its own supported authentication and service implementation.')
p('The previous Teams error text does not, by itself, identify a failed program integration. It could not establish a cause for the Microsoft error. A repair claim would require reproducing the failure and verifying the relevant action in the correct account. No such claim is made here.')
h('Recommendation')
p('Continue using clear destination links while the bridge is off. Label the owner view honestly. Do not present reminders, membership synchronization, or document transfer as operating until a complete action has been verified in both systems.')
ref('S09 S10 S12')

page('12 Evaluation maintenance and release')
p('Evaluation is a workflow in its own right. The owner can request defined suites; the runner uses isolated test storage and records results. Five named suites cover Ask, cultural intelligence, related quality, independent practice, and the program mindset. Manual checks remain manual. Merely having cases on disk is not evidence that they currently pass.')
table(['Trigger','Sequence and result'],[
('Owner evaluation request','Authenticate, select a suite, run its cases in the intended test mode, record results, and review failures. Model comparison remains a disabled catalog capability.'),
('Code push or pull request','The GitHub workflow checks a clean copy, installs locked dependencies, prepares PostgreSQL test tools, then runs environment, source, type, lint, test, and build checks.'),
('Manual verification','The same verification command can be run by the developer. The workflow checks that verification does not modify tracked files.'),
('Deployment build','The Vercel build script supplies deployment specific checks. Successful local rendering does not prove a deployment or a live service call.'),
('Health request','A health endpoint supports readiness inspection. Its output is a diagnostic indicator, not proof that every user journey works.'),
('Scheduled maintenance','The configured cycle can perform retention and review duties once actually scheduled and authorized. Local server availability alone is insufficient.')],[1.43,5.57])
h('What was verified for this report')
p('The initial inspection was followed by targeted corrections at the owner’s request. Automated regression checks and TypeScript validation were run for those changes; the correction record at the end explains their scope. Saved page editing notes supply earlier observed local evidence. Live provider calls, an unattended schedule, and end to end Teams actions were not tested here.')
h('Release evidence worth keeping')
p('For each major workflow, retain its configuration state, a successful scenario, a failure scenario, the resulting saved state, and a recovery test. Keep this evidence in the consultant area. Staff should see useful service language, not the internal mechanics of development or testing.')
ref('S13 S15 S16')

page('13 Priorities for completing the ecosystem')
p('The completion sequence below remains open except for the targeted control and review corrections recorded in section 17. It preserves the ambitious program direction without mistaking scaffolding for an operating service. Larger integrations and resource activation are not complete.')
table(['Priority','Action and acceptance evidence'],[
('1 Resource visibility','Reconcile source records, usable assets, released resources, and visible results. Keep seeds. Show representative recovered courses and briefs in the selected staff scope and in Ask.'),
('2 Intelligent assistance','Connect the chosen supported provider, explicitly enable research, and test real questions, absent sources, privacy refusal, provider failure, and cost reporting. Do not imply fresh research during fallback.'),
('3 Consistent owner control','Test flags, agent disablement, authority levels, and the stop control across direct cycle calls and generative paths. Document the precise limits of undo.'),
('4 Durable service workflows','Complete consultation activation evidence before accepting requests. Prove tracking, correction, withdrawal, retention, failed delivery handling, and restoration arrangements.'),
('5 Collaboration adoption','Separate owner editing from member participation. Prove member permissions before inviting a team into the native workspace. Keep personal Teams navigation distinct from automation.'),
('6 Role specific learning','Validate organizational roles with responsible leaders, connect them to resources and Ask, and verify relevance without inferring identity or assessing staff beliefs.'),
('7 Sustainable growth','Measure large collection ingestion and graph retrieval, review content ownership, test model upgrades, and repeat program review over time.')],[1.48,5.52])
h('The program standard')
p('The intended result remains a continuous equity program, not a collection of disconnected tools. Staff should encounter resources relevant to their work, choose an accessible route into learning, and see meaningful follow through. The consultant should be able to direct and stop the system without becoming the manual operator for every interaction.')
p('That standard is compatible with experimentation. The distinction is that each experiment needs an actual working path, an observable result, and a defined boundary—not simply a promising agent name.')

# Complete catalog extracted from the inspected source without treating declarations as execution.
catalog=(ROOT/'lib/intelligence/tools/catalog.ts').read_text(encoding='utf-8')
pattern=r't\("([^"\n]+)", "([^"\n]+)", "([^"\n]+)", "([^"\n]+)", "([^"\n]+)", "([^"\n]+)", (true|false), "([^"\n]*)"(?:, (true|false))?\)'
entries=re.findall(pattern,catalog)
assert len(entries)==len(re.findall(r'^\s+t\(',catalog,re.M)), 'Catalog parser missed an entry'
for part,start in enumerate(range(0,len(entries),18),1):
    page('14 Tool register '+str(part))
    p('Every named tool is listed below. “On” means enabled in the catalog, not proven active. Authority is the minimum declared level. Some tools are helpers or declarations rather than independent services; the preceding sections explain the actual limits.')
    rows=[]
    for name,fam,desc,perm,auth,phase,side,safety,explicit in entries[start:start+18]:
        on=explicit=='true' if explicit else phase=='mvp'
        rows.append((name,desc,('On' if on else 'Off')+'  '+auth))
    table(['Registered tool','Declared purpose','Catalog'],rows,[2.07,4.17,.76],9.5)
    ref('S06  Full catalog of '+str(len(entries))+' tools. The catalog is an inventory, not a completion claim.')

page('15 Application entry points')
p('These are the current service entry points found under app/api. Most are triggered by a person using the application, not by a schedule. Names are included for maintenance and independent review; they are not proposed staff facing labels.')
table(['Route under app api','Workflow'],[
('ask','Question handling and optional research.'),
('intake and intake/[id]','Preparation or submission, and protected requester actions.'),
('cron/orchestrate','Secret protected scheduled review cycle.'),
('health','Readiness and diagnostic response.'),
('consultant/login and logout','Owner session entry and exit.'),
('consultant/orchestrator','Owner cycle, policy, proposal, and undo controls.'),
('consultant/research','Owner research settings and related state.'),
('consultant/evals','Owner initiated evaluation.'),
('consultant/review','Consultant review work.'),
('consultant/queue/[id]','Owner consultation queue updates.'),
('consultant/one-dsd-team','Owner controlled native collaboration changes.'),
('consultant/content/[surfaceId]','Registered page draft, review, publication, and restoration actions.'),
('consultant/page-copy/[surface]','Separate page copy path; not universal resource or layout editing.'),
('consultant/resources/[id]/draft and release','Resource draft and publication lifecycle.')],[3.05,3.95],10)
p('Owner authentication is separate from personal Microsoft sign in. The current protected application paths use an owner session; they do not establish a full staff, co lead, equity director, and assistant commissioner permission hierarchy. Names and roles in a data model do not supply that authorization by themselves.')
ref('S11 S12 S13 S16')

page('16 Evidence and review boundaries')
p('All references below are relative to the inspected one-dhs-pac repository in the Downloads project folder. They identify the implementation evidence without exposing credentials. Configuration findings apply to the inspected local setup, not automatically to Vercel or another developer machine.')
sources=[
('S01','lib/product/federation.ts and request-context.ts; scoped program context.'),
('S02','lib/content/learning-hub.ts, learning-catalog.ts, work-learning.ts; lib/product/learning.ts.'),
('S03','lib/product/workforce-map.ts; workforce views and editable content definitions.'),
('S04','lib/intelligence/orchestrator.ts; agents/ask.ts; retrieval/search.ts and staff-search.ts.'),
('S05','lib/intelligence/research/index.ts, governance.ts, providers.ts; local configuration indicators.'),
('S06','lib/intelligence/registry/agents.ts; tools/catalog.ts and runtime.ts; registry/flags.ts and models.ts.'),
('S07','lib/intelligence/policy.ts; providers/index.ts; agents/cycle.ts.'),
('S08','app/api/cron/orchestrate/route.ts; consultant/orchestrator/route.ts; vercel.json.'),
('S09','lib/collaboration/schema.ts, store.ts, view.ts, seed.ts, microsoft-boundary.ts.'),
('S10','One DSD Team and Amplify pages and content; owner program direction in this conversation.'),
('S11','lib/intelligence/consult/; agents/intake.ts; memory/consultation-cas.ts; intake routes.'),
('S12','Content editing routes and registry; migration 0018; local selected editing areas.'),
('S13','docs/Saved page editing verification.txt; docs/Local consultant access verification.txt; auth routes.'),
('S14','scripts/corpus/; package.json ingestion commands; source and canonical projection documentation.'),
('S15','lib/intelligence/eval/runner.ts; consultant/evals route; evaluation suites.'),
('S16','.github/workflows/verify.yml; scripts/vercel-build.mjs; package.json; app/api route inventory.')]
table(['Reference','Evidence'],sources,[.8,6.2],9.5)
p('No live deployment was certified. Historical resource totals are not a current release count. No private staff conversation, credential, or consultation record is reproduced. Diagrams describe code paths where stated and clearly identify human led or proposed connections. This is a program wide implementation review, not a security certification or proof that every branch has passed testing.')

page('17 Corrections and remaining work')
p('After requesting this document, the owner authorized corrections to missing and incomplete behavior. The following targeted repairs were implemented locally. They do not represent completion of the entire program.')
table(['Correction','Verification provided'],[
('Generated proposal authority','A model response cannot save new proposals below A5. The save also checks policy after generation, so a newly applied stop blocks it.'),
('Disabled orchestrator','A lower authority cycle does not bypass orchestrator disablement. Required privacy retention and security cleanup remain available.'),
('Specialist controls','Brief review and accessibility scanning respect their assigned agent settings. A2 accessibility work is denied at A0.'),
('Flags and scope','The shared runner enforces agent flags, role and environment scope, and the separate stale flag writing switch.'),
('Model readiness','A candidate model is not reported active solely because its credential and pilot flag exist.'),
('Published collection review','Stale detection and accessibility review use the selected published collection rather than only seed definitions.'),
('Accurate recovery description','The triage catalog now states that undo clears eligible priority pins, not consultation status history.')],[1.53,5.47])
p('Verification: eleven new regression tests cover these controls and published collection selection. A focused run passed 44 tests; a broader run passed 165 tests covering owner access, publication, privacy, consultation, and retention. These sets overlap and should not be added together. TypeScript validation passed. Provider responses were simulated; no paid call or live integration is claimed.')
h('Still open')
p('Full resource publication and reconciliation, graph powered retrieval, role specific adaptive learning, the OpenAI adapter, live research credentials and activation, consultation readiness, member collaboration permissions, Teams transport, notifications, and verified scheduling remain separate completion work. Nothing in this document marks those items finished.')
p('The staff interface, DHS logo, content source, deployment, and account settings were not changed by this repair pass. No database migration was applied. Test records used isolated test storage.')
ref('tests/workflow-control-completion.test.ts and the implementation references for sections 5 6 and 12')

doc.save(OUT)
print(str(OUT))
print('Catalog entries:',len(entries))
